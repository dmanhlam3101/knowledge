# -*- coding: utf-8 -*-
"""Sinh file kich ban kiem thu .xlsx tu khuon 08-kich-ban-kiem-thu.xlsx.

Cach dung:
    python "AI Analysis/ba/_tools/gen_kich_ban_kiem_thu.py" <spec.json> [-o <out.xlsx>]

Spec JSON (UTF-8) - xem vi-du-kich-ban.json:
{
  "ma": "NV_08",
  "ten": "Them han xu ly o Danh sach nguoi cung nhan",
  "out": "AI Analysis/ba/yeu-cau/.../testcase/NV_08.xlsx",
  "sheets": [
    {"kenh": "Web", "precondition": "...", "rows": [
        {"loai": "G", "text": "..."},
        {"loai": "S", "text": "..."},
        {"loai": "N", "text": "N/A - ly do"},
        {"loai": "D", "muc_dich": "...", "buoc": "...", "ket_qua": "...", "ghi_chu": "..."}
    ]}
  ]
}

Luat giu nguyen tu khuon:
- Cot A sinh ma tu dong bang cong thuc, lay tien to tu o D3.
- Cot Q suy ket qua tu cac cot trinh duyet.
- D4..D8 la cong thuc dem, script chi noi dai tham chieu theo so dong moi.
- Cot ket qua (E..P) luon de TRONG: test chua chay thi khong dien.
Khong co LibreOffice thi file van dung, Excel tu tinh cong thuc khi mo.
"""
from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
from copy import copy

import openpyxl
from openpyxl.worksheet.cell_range import CellRange

HERE = os.path.dirname(os.path.abspath(__file__))
BA_DIR = os.path.dirname(HERE)
REPO = os.path.dirname(os.path.dirname(BA_DIR))
TEMPLATE = os.path.join(BA_DIR, "templates", "08-kich-ban-kiem-thu.xlsx")

SHEET_WEB_SRC = "YC02_Web"
SHEET_APP_SRC = "YC02_App"
LAST_COL = 19  # A..S
ROW_TITLE = 12
ROW_PRE = 13
ROW_BODY = 14

A_TMPL = '=IF(AND(C{r}="",C{r}=""),"",$D$3&"_"&ROW()-12-COUNTBLANK($C$13:C{r}))'

Q_TMPL = (
    '=IF(OR(IF(G{r}="",IF(F{r}="",IF(E{r}="","",E{r}),F{r}),G{r})="F",'
    'IF(J{r}="",IF(I{r}="",IF(H{r}="","",H{r}),I{r}),J{r})="F",'
    'IF(M{r}="",IF(L{r}="",IF(K{r}="","",K{r}),L{r}),M{r})="F",'
    'IF(P{r}="",IF(O{r}="",IF(N{r}="","",N{r}),O{r}),P{r})="F")=TRUE,"F",'
    'IF(OR(IF(G{r}="",IF(F{r}="",IF(E{r}="","",E{r}),F{r}),G{r})="PE",'
    'IF(J{r}="",IF(I{r}="",IF(H{r}="","",H{r}),I{r}),J{r})="PE",'
    'IF(M{r}="",IF(L{r}="",IF(K{r}="","",K{r}),L{r}),M{r})="PE",'
    'IF(P{r}="",IF(O{r}="",IF(N{r}="","",N{r}),O{r}),P{r})="PE")=TRUE,"PE",'
    'IF(AND(IF(G{r}="",IF(F{r}="",IF(E{r}="","",E{r}),F{r}),G{r})="",'
    'IF(J{r}="",IF(I{r}="",IF(H{r}="","",H{r}),I{r}),J{r})="",'
    'IF(M{r}="",IF(L{r}="",IF(K{r}="","",K{r}),L{r}),M{r})="",'
    'IF(P{r}="",IF(O{r}="",IF(N{r}="","",N{r}),O{r}),P{r})="")=TRUE,"","P")))'
)


def build_sheet(ws, sheet_spec, doc_name, code_prefix):
    ws["D2"] = doc_name
    ws["D3"] = code_prefix
    ws[("B%d" % ROW_TITLE)] = "%s_%s\n" % (doc_name, sheet_spec["kenh"])
    ws[("B%d" % ROW_PRE)] = sheet_spec["precondition"]
    if sheet_spec.get("cao_dong_tien_dieu_kien"):
        ws.row_dimensions[ROW_PRE].height = sheet_spec["cao_dong_tien_dieu_kien"]

    sty_group = [copy(ws.cell(row=14, column=c)._style) for c in range(1, LAST_COL + 1)]
    sty_sub = [copy(ws.cell(row=17, column=c)._style) for c in range(1, LAST_COL + 1)]
    sty_data = [copy(ws.cell(row=18, column=c)._style) for c in range(1, LAST_COL + 1)]

    for rng in [str(r) for r in ws.merged_cells.ranges]:
        if CellRange(rng).min_row >= ROW_BODY:
            ws.unmerge_cells(rng)
    if ws.max_row >= ROW_BODY:
        ws.delete_rows(ROW_BODY, ws.max_row - ROW_BODY + 1)

    r = ROW_BODY
    n_tc = 0
    for item in sheet_spec["rows"]:
        kind = item["loai"].upper()
        if kind in ("G", "S", "N"):
            sty = sty_group if kind == "G" else (sty_sub if kind == "S" else sty_data)
            for c in range(1, LAST_COL + 1):
                ws.cell(row=r, column=c)._style = copy(sty[c - 1])
            ws.cell(row=r, column=1).value = A_TMPL.format(r=r)
            ws.cell(row=r, column=2).value = item["text"]
            ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=17)
            ws.row_dimensions[r].height = 16.15 if kind in ("G", "S") else None
        elif kind == "D":
            for c in range(1, LAST_COL + 1):
                ws.cell(row=r, column=c)._style = copy(sty_data[c - 1])
            ws.cell(row=r, column=1).value = A_TMPL.format(r=r)
            ws.cell(row=r, column=2).value = item["muc_dich"]
            ws.cell(row=r, column=3).value = item["buoc"]
            ws.cell(row=r, column=4).value = item["ket_qua"]
            ws.cell(row=r, column=17).value = Q_TMPL.format(r=r)
            if item.get("ghi_chu"):
                ws.cell(row=r, column=19).value = item["ghi_chu"]
            ws.row_dimensions[r].height = None
            n_tc += 1
        else:
            raise ValueError("loai khong hop le: %r (chi nhan G, S, N, D)" % kind)
        r += 1

    last = r - 1
    ws["D4"] = '=COUNTIF($E$13:$E${0},"P")'.format(last)
    ws["D5"] = '=COUNTIF($E$13:$E${0},"F")'.format(last)
    ws["D6"] = '=COUNTIF($Q$13:$Q${0},"PE")'.format(last)
    ws["D7"] = "=D8-D4-D5-D6"
    ws["D8"] = "=COUNTA($Q$13:$Q${0})".format(last)
    return last, n_tc


def build_summary(wb, sheet_names):
    ws = wb["Tổng hợp"]
    for r in range(4, 12):
        for c in range(1, 11):
            ws.cell(row=r, column=c).value = None
    for i, name in enumerate(sheet_names):
        r = 4 + i
        ws.cell(row=r, column=1).value = i + 1
        ws.cell(row=r, column=2).value = "='{0}'!D2&\" - {1}\"".format(
            name, name.rsplit("_", 1)[-1])
        for col, src in [(3, "D4"), (4, "D5"), (5, "D6"), (6, "D7"), (7, "D8")]:
            ws.cell(row=r, column=col).value = "='{0}'!{1}".format(name, src)
        ws.cell(row=r, column=8).value = "=IFERROR(C{0}/G{0},0)".format(r)
        ws.cell(row=r, column=9).value = "=IFERROR(D{0}/G{0},0)".format(r)
        ws.cell(row=r, column=10).value = "=IFERROR((C{0}+D{0})/G{0},0)".format(r)
    ws["H12"] = "=IFERROR(C12/G12,0)"
    ws["I12"] = "=IFERROR(D12/G12,0)"
    ws["J12"] = "=IFERROR((C12+D12)/G12,0)"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("-o", "--out")
    ap.add_argument("--template", default=TEMPLATE)
    args = ap.parse_args()

    with open(args.spec, encoding="utf-8") as fh:
        spec = json.load(fh)

    out = args.out or spec.get("out")
    if not out:
        raise SystemExit("Thieu duong dan ra: dat 'out' trong spec hoac dung -o")
    if not os.path.isabs(out):
        out = os.path.join(REPO, out)
    os.makedirs(os.path.dirname(out), exist_ok=True)

    shutil.copy(args.template, out)
    wb = openpyxl.load_workbook(out)

    src_sheets = [SHEET_WEB_SRC, SHEET_APP_SRC]
    doc_name = "%s_%s" % (spec["ma"], spec["ten"])
    made, report = [], []

    for i, sh in enumerate(spec["sheets"]):
        if i >= len(src_sheets):
            raise SystemExit("Khuon chi co 2 sheet kich ban (Web, App)")
        ws = wb[src_sheets[i]]
        last, n_tc = build_sheet(ws, sh, doc_name, spec["ma"])
        new_name = "%s_%s" % (spec["ma"], sh["kenh"])
        ws.title = new_name
        made.append(new_name)
        report.append((new_name, n_tc, last))

    for leftover in src_sheets[len(spec["sheets"]):]:
        del wb[leftover]

    build_summary(wb, made)
    wb["Trang bìa"]["A11"] = "%s - %s" % (spec["ma"], spec["ten"])
    wb.save(out)

    print("DA TAO:", out)
    for name, n_tc, last in report:
        print("  %-16s %3d testcase  (dong cuoi %d)" % (name, n_tc, last))
    print("Cot ket qua de trong - Excel se tu tinh cac o dem khi mo file.")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
